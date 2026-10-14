"""The report: wording, the data it is built from, its three file formats, and its draft-to-approved lifecycle."""

import csv
import io
import json
from unittest import mock

import pytest
from openpyxl import load_workbook

from conftest import ids
from fairprice import agent, db, report, server
from fairprice.storage import storage

HILUX = "SELECT id FROM lines WHERE description LIKE 'Supply of Toyota Hilux double cabin pickup unit%' ORDER BY id"
BASE = {"currency": "NGN", "year": 2025, "country": "NG"}


def checked(run, line, candidates, rejected=None):
    """A line searched and checked through the run's own tools (the search step stands in for find_comparables)."""
    db.add_step(run, line, "search", {"keywords": ["hilux"], "country": "NG", "matches": 11, "ids": candidates})
    return server.check_price_comparability(run, candidates, rejected, line)


def package_run(title="pkg", budget=None):
    """Two lines, three bidders: one line with a band, one with too few records, and one bidder missing a line."""
    lines = [{"line": "1", "item": "Toyota Hilux double cabin pickup", "quantity": 2,
              "prices": {"A": 95e6, "B": 300e6, "C": 40e6}, **({"budget": budget} if budget else {})},
             {"line": "2", "item": "Laptop computer", "quantity": 10, "prices": {"A": 2e6, "B": 2.1e6}}]
    checks = [{"check": "arithmetic_error", "bidder": "B", "line": "1", "where": "sheet “B”, row 2",
               "data": {"quantity": 2, "unit_price": 300e6, "total": 700e6}},
              {"check": "missing_line", "bidder": "C", "line": "2", "where": None, "data": {"item": "Laptop computer"}}]
    run = db.start_run("package", title, BASE, lines, ["A", "B", "C"], checks)
    bundle = ids("SELECT id FROM lines WHERE description LIKE '%Prado%'")
    checked(run, "1", ids(HILUX) + bundle, {str(bundle[0]): "Challenger: bundle"})
    server.challenge_review(run, "STANDS", "fine")
    checked(run, "2", ids("SELECT id FROM lines WHERE description LIKE 'Laptop computer%'"))
    return run


@pytest.fixture
def package():
    run = package_run()
    record, lines, states = server.review(run)
    return lines, states, record.bidders, record.checks


def test_money_and_ranges():
    assert report.money(101990442, "NGN") == "₦101,990,442"
    assert [report.money(v, "NGN", True) for v in (1.25e6, 300e6, 2e9, 950)] == ["₦1.25M", "₦300M", "₦2B", "₦950"]
    assert report.money(5, "USD") == "USD 5" and report.money(None, "NGN") == ""
    assert report.ranges([5, 1, 2, 3]) == "1–3, 5" and report.ranges([]) == ""


def test_plain_reasons():
    assert report.plain_reason("Challenger: a bundle") == "Second check: a bundle"
    assert report.plain_reason("invalid_date:2001, multi_item_contract") == (
        "the date is missing or a placeholder (2001); the contract covers several different items, so it has no price per item")
    assert report.plain_reason("invalid_date:None") == "the date is missing or a placeholder"
    assert report.plain_reason("different model") == "different model"


def test_cells_judge_every_quote(package):
    lines, states, bidders, _ = package
    cells = report.cells(lines, {**states, "3": {"error": "boom"}}, bidders)
    assert cells[("1", "A")]["status"] == "within_band"
    assert cells[("1", "B")]["status"] == "above_band" and cells[("1", "C")]["status"] == "below_band"
    assert cells[("2", "A")]["status"] == "no_comparison" and cells[("2", "C")]["status"] == "not_priced"
    failed = report.cells([{"line": "3", "item": "x", "quantity": 1, "prices": {"A": 1}}], {"3": {"error": "boom"}}, ["A"])
    assert failed[("3", "A")]["status"] == "failed"


def test_check_sentences():
    where = {"bidder": "B", "line": "2", "where": "row 3"}
    assert "does not add up" in report.check_text({**where, "check": "arithmetic_error",
                                                   "data": {"quantity": 2, "unit_price": 5, "total": 11}}, "NGN")
    assert "asks for 2" in report.check_text({**where, "check": "quantity_changed", "data": {"quoted": 1, "requested": 2}}, "NGN")
    assert report.check_text({**where, "check": "missing_line", "data": {"item": "Laptop"}}, "NGN") == \
        "B did not price line 2 (Laptop)."  # facts only: no country's law is quoted
    assert "A and C quoted exactly the same" in report.check_text(
        {"check": "identical_prices", "bidder": "A, C", "line": "1, 2", "where": None, "data": {"same": 2, "shared": 2}}, "NGN")
    assert report.check_text({**where, "check": "something_new", "data": {}}, "NGN") == "B, line 2 (row 3): something_new."


def test_terminal_summary(package):
    lines, states, bidders, checks = package
    text = report.summary(lines, states, bidders, checks, BASE)
    assert "Line 1 · Toyota Hilux double cabin pickup: typical past price" in text
    assert "Well above past prices (" in text and "Not priced" in text and "Problems in the bid files: 2" in text


def test_report_data_cites_its_records_and_never_ranks(package):
    lines, states, bidders, checks = package
    data = report.build("pkg.xlsx", lines, states, bidders, checks, BASE, "run1", "ollama:gemma4")
    assert data["report"]["status"] == "draft" and data["report"]["prepared_by"] == "FairPrice (ollama:gemma4)"
    assert [p["bidder"] for p in data["prices"][:3]] == bidders  # the file's order, not ranked
    assert data["prices"][0]["sources"] and data["prices"][2]["result"] == "Well below past prices"
    assert len({s["record_id"] for s in data["sources"]}) == len(data["sources"]) >= 8
    assert data["report"]["bid_file_checks"][1] == "C did not price line 2 (Laptop computer)."
    assert "budget_per_item" not in data["prices"][0]  # no budget in this BoQ, so no budget columns
    assert not any("rank" in k or "total" in k for k in data["prices"][0])
    single = report.build("x", *report.single(states["1"]), None, states["1"]["request"], "run2")
    assert single["report"]["bid_file_checks"] is None and single["prices"][0]["bidder"] == "Price under review"


def test_status_line_and_headings(package):
    lines, states, bidders, checks = package
    data = report.build("pkg", lines, states, bidders, checks, BASE, "run1")
    assert report.status_line(data) == report.DRAFT_STATUS
    data["report"].update(status="approved", approved_by="Ada", approved_at="2026-10-12T10:00:00+01:00")
    assert report.status_line(data) == "Approved by Ada on 12 October 2026 at 10:00"
    prices, srcs, docs = report.headings(data)
    assert docs == []  # no checklist in this package
    assert prices[0]["Quoted price per item (NGN)"] == 95000000 and isinstance(prices[0]["Quantity requested"], int)
    assert prices[0]["Sources (Ref numbers)"].startswith("1") and srcs[0]["FairPrice audit ID"].startswith("L")


def test_three_file_formats(package):
    lines, states, bidders, checks = package
    data = report.build("pkg", lines, states, bidders, checks, BASE, "run1")
    xlsx = report.files(data, "R", "xlsx")
    wb = load_workbook(io.BytesIO(xlsx["R.xlsx"]))
    assert wb.sheetnames == ["Prices", "Sources"] and wb["Prices"].oddHeader.center.text == report.DRAFT_STATUS
    assert wb["Prices"]["H3"].fill.fgColor.rgb.endswith("F8D7D7")  # Bidder B, well above: red
    assert wb["Prices"].cell(wb["Prices"].max_row, 1).value == report.DRAFT_STATUS
    out = report.files(data, "R", "csv")
    assert set(out) == {"R-prices.csv", "R-sources.csv"} and out["R-prices.csv"].startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(out["R-prices.csv"].decode("utf-8-sig"))))
    assert rows[0]["Status"] == report.DRAFT_STATUS and rows[0]["Bidder"] == "A"
    assert json.loads(report.files(data, "R", "json")["R.json"]) == data
    empty = {**data, "sources": []}
    assert report.files(empty, "E", "csv")["E-sources.csv"].decode("utf-8-sig").startswith("Status")


def test_draft_approve_retrieve_and_reject(monkeypatch):
    for fmt in report.FORMATS:
        monkeypatch.setenv("FAIRPRICE_REPORT_FORMAT", fmt)
        run = package_run("Bid Package!")
        asked = server.write_price_review(run)  # no person, no approval: a draft and a request
        assert asked["status"] == "needs_approval" and asked["draft"] == f"DRAFT-bid-package-{run}"
        assert storage().load(asked["files"][0].rsplit("/", 1)[-1])
        db.decide(asked["request"], True, "Ada Obi")
        saved = server.write_price_review(run, asked["request"])
        assert saved["saved"].endswith(f"bid-package-{run}") and saved["sources"] >= 8
        got = server.get_price_review(run)
        assert got["report"]["approved_by"] == "Ada Obi" and got["files"]
        assert got["lines"][0]["second_check"]["verdict"] == "STANDS"
        assert any(r["run_id"] == run for r in server.get_price_review()["reports"])
        with pytest.raises(Exception):
            storage().load(asked["files"][0].rsplit("/", 1)[-1])
    run = package_run("rejected")
    asked = server.write_price_review(run)
    server.discard_review(run)  # the person rejected it
    db.decide(asked["request"], True, "Ada")  # too late: the request was cancelled with its draft
    assert "not an unused human approval" in server.write_price_review(run, asked["request"])["error"]
    assert "No approved report" in server.get_price_review("nothing")["error"]
    assert "No review" in server.write_price_review("nothing")["error"]


def test_approval_refusals():
    run = package_run("refusals")
    assert "not an unused human approval" in server.write_price_review(run, "made-up")["error"]
    other = db.grant("write_price_review", f"{package_run('other')}@1", "Ada")
    assert "not an unused human approval" in server.write_price_review(run, other)["error"]
    first = server.write_price_review(run)["request"]
    second = server.write_price_review(run)["request"]  # a new draft cancels the old one and its request
    db.decide(first, True, "Ada")
    assert "not an unused human approval" in server.write_price_review(run, first)["error"]
    db.decide(second, True, "Ada")
    checked(run, "2", ids("SELECT id FROM lines WHERE description LIKE 'Laptop computer%'"))  # the review changed
    assert "not an unused human approval" in server.write_price_review(run, second)["error"]
    empty = db.start_run("item", "nothing yet", BASE, [{"line": "1", "item": "x", "quantity": None, "prices": {"P": 1}}], ["P"])
    assert "nothing to report" in server.write_price_review(empty)["error"]


def test_opening_the_draft_only_when_asked(monkeypatch):
    with mock.patch("subprocess.run") as run:
        agent.open_draft("R.xlsx")
        run.assert_not_called()
        monkeypatch.setenv("FAIRPRICE_OPEN_DRAFT", "1")
        agent.open_draft("R.xlsx")
        run.assert_called_once()
    monkeypatch.setattr("sys.platform", "win32")
    with mock.patch("os.startfile", create=True) as start:
        agent.open_draft("R.xlsx")
    start.assert_called_once()


def test_budget_columns_when_the_boq_has_budgets():
    record, lines, states = server.review(package_run("budgets", budget=100e6))
    data = report.build("pkg", lines, states, record.bidders, record.checks, BASE, "b1")
    a, b = data["prices"][0], data["prices"][1]
    assert (a["budget_per_item"], a["times_budget"], b["times_budget"]) == (100e6, 0.95, 3.0)
    assert data["prices"][3]["budget_per_item"] is None  # line 2 has no budget
    assert "Well above past prices (" in report.summary(lines, states, record.bidders, [], BASE) and \
        "3× the budget" in report.summary(lines, states, record.bidders, [], BASE)
    ws = load_workbook(io.BytesIO(report.files(data, "B", "xlsx")["B.xlsx"]))["Prices"]
    assert ws["I1"].value == "Budget per item (NGN)" and ws["H1"].value == "Result"
    assert ws["H3"].fill.fgColor.rgb.endswith("F8D7D7")  # the result colour stays on the Result column


def test_the_documents_checklist_in_the_report():
    documents = [{"ref": "1", "document": "Tender security", "rule": "Submitted"},
                 {"ref": "2", "document": "Tax compliance certificate", "rule": "Valid on 2025-10-15"}]
    run = package_run("documents")
    record, lines, states = server.review(run)
    checks = [{"check": "document_missing", "bidder": "C", "line": None, "where": "sheet “Documents”, row 2, C",
               "data": {"ref": "1", "document": "Tender security"}},
              {"check": "document_expired", "bidder": "B", "line": None, "where": "sheet “Documents”, row 3, B",
               "data": {"ref": "2", "document": "Tax compliance certificate", "valid_until": "2025-09-30",
                        "needed": "2025-10-15"}}]
    base = {**BASE, "documents": documents}
    data = report.build("pkg", lines, states, record.bidders, checks, base, "d1")
    assert [p["documents"] for p in data["prices"][:3]] == ["All required documents", "1 document problem",
                                                            "1 document problem"]
    assert data["documents"][0]["results"] == {"A": "OK", "B": "OK", "C": "Not submitted"}
    assert data["documents"][1]["results"]["B"] == "Expired"
    assert "line 1" not in " ".join(data["prices"][0]["problems"])  # document problems are not price problems
    text = report.summary(lines, states, record.bidders, checks, base)
    assert "Documents (2 required): A: All required documents; B: 1 document problem" in text
    assert "valid until 2025-09-30, but the tender requires it to be valid on 2025-10-15" in text
    wb = load_workbook(io.BytesIO(report.files(data, "D", "xlsx")["D.xlsx"]))
    assert wb.sheetnames == ["Prices", "Sources", "Documents"]
    assert [c.value for c in wb["Documents"][1]] == ["Ref", "Document", "Rule", "A", "B", "C"]
    assert wb["Documents"]["F2"].fill.fgColor.rgb.endswith("F8D7D7") and wb["Documents"]["D2"].fill.fgColor.rgb in (None, "00000000")
    assert "D-documents.csv" in report.files(data, "D", "csv")
