"""Bid packages: reading xlsx, csv and json against the schema, refusals with locations, the checks, and a full run."""

import argparse
import asyncio
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from conftest import ScriptedModel, fake_session
from fairprice import batch, db, server

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


@pytest.mark.parametrize("ext", ["xlsx", "csv", "json"])
def test_templates_read_the_same_and_find_the_planted_problems(ext):
    boq, bidders = batch.read_package(EXAMPLES / f"bid-package.{ext}")
    assert [r["line"] for r in boq] == ["1", "2", "3"] and list(bidders) == ["Bidder A", "Bidder B", "Bidder C"]
    found = {(c["check"], c["bidder"], c["line"]) for c in batch.file_checks(boq, bidders)}
    assert found == {("arithmetic_error", "Bidder B", "2"), ("quantity_changed", "Bidder C", "1"),
                     ("missing_line", "Bidder C", "3"), ("identical_prices", "Bidder A, Bidder C", "1, 2")}
    assert "Bidder C: 2 of 3 lines priced (not priced: 3)" in batch.describe(f"x.{ext}", boq, bidders)
    lines = batch.package_lines(boq, bidders)
    assert lines[2]["prices"] == {"Bidder A": 1300000.0, "Bidder B": 1250000.0}


def refused(path):
    with pytest.raises(batch.PackageError) as e:
        batch.load(path)
    return "\n".join(e.value.problems)


def test_refusals_name_the_problem_and_where_it_is(tmp_path):
    wb = load_workbook(EXAMPLES / "bid-package.xlsx")
    wb["Bidder B"]["E1"] = "Price per piece"
    wb["Bidder A"]["G2"] = "stray"  # a value under no heading
    wb["Bidder C"]["C1"] = "Rate"  # a second unit price column
    wb.save(tmp_path / "a.xlsx")
    problems = refused(tmp_path / "a.xlsx")
    assert "sheet “Bidder B”, row 1: unknown column(s) Price per piece" in problems
    assert "missing column(s) Unit Price" in problems and "the columns are Item No, Description" in problems
    assert "sheet “Bidder A”, row 2: a value in a column with no name" in problems
    assert "sheet “Bidder C”, row 1: more than one column for Unit Price" in problems
    wb = load_workbook(EXAMPLES / "bid-package.xlsx")
    wb.remove(wb["BoQ"])
    wb.create_sheet("Empty")
    wb.save(tmp_path / "b.xlsx")
    problems = refused(tmp_path / "b.xlsx")
    assert 'no sheet named "BoQ"' in problems and "sheet “Empty”: no rows" in problems


def test_csv_refusals(tmp_path):
    head = "Tenderer,Item No.,Description,Qty,Rate in KES (VAT inclusive),Total Amount (KES)\n"  # KRA's headings
    cases = {"BoQ,1,Laptop,2,,\nBidder A,1,,2,NGN 5,\n": "row 3, Rate in KES (VAT inclusive): 'NGN 5' is not a number",
             "BoQ,1,Laptop,2,5,\n": "a BoQ row has no Unit Price",
             ",1,Laptop,2,,\n": "no bidder (write BoQ",
             "Bidder A,1,,2,5,\n": "no BoQ rows",
             "BoQ,1,Laptop,2,,\nBoQ,1,Laptop,2,,\nBidder A,1,,2,5,\nBidder A,1,,2,6,\nBidder A,7,,2,6,\n":
                 "BoQ line 1 appears more than once"}
    for i, (rows, expected) in enumerate(cases.items()):
        (tmp_path / f"{i}.csv").write_text(head + rows)
        assert expected in refused(tmp_path / f"{i}.csv")
    (tmp_path / "dup.csv").write_text(head + "BoQ,1,Laptop,2,,\nBidder A,1,,2,5,\nBidder A,1,,2,6,\nBidder A,7,,2,6,\n")
    problems = refused(tmp_path / "dup.csv")
    assert "line 1 is priced twice" in problems and "line 7 is not in the BoQ" in problems


def test_json_refusals(tmp_path):
    pkg = json.loads((EXAMPLES / "bid-package.json").read_text())
    pkg["bids"].append(dict(pkg["bids"][0]))
    (tmp_path / "dup.json").write_text(json.dumps(pkg))
    assert "bidder 'Bidder A' appears more than once" in refused(tmp_path / "dup.json")
    del pkg["boq"]
    (tmp_path / "noboq.json").write_text(json.dumps(pkg))
    assert "the package: 'boq' is a required property" in refused(tmp_path / "noboq.json")
    (tmp_path / "broken.json").write_text("{not json")
    assert "not valid JSON" in refused(tmp_path / "broken.json")
    (tmp_path / "x.txt").write_text("")
    assert "use .xlsx, .csv or .json" in refused(tmp_path / "x.txt")
    assert "file not found" in refused(tmp_path / "nowhere.xlsx")
    with pytest.raises(batch.PackageError) as e:
        batch.load(tmp_path / "broken.json")
    assert "was not read: it does not match the bid package format" in str(e.value)


def test_ppra_headings_and_budgets():
    for heading, field in (("Item No.", "line"), ("Rate in KES (VAT inclusive)", "unit_price"), ("QTY", "quantity"),
                           ("Total Amount (KES)", "total"), ("Budget per unit (KES)", "budget"), ("Tenderer", "bidder")):
        assert batch.FIELD[batch.plain(heading)] == field
    assert batch.plain("price in words") not in batch.FIELD
    boq, bidders = batch.read_package(EXAMPLES / "bid-package.xlsx")
    assert boq[1]["unit"] == "No" and boq[1]["budget"] == 2400000
    assert batch.package_lines(boq, bidders)[1]["budget"] == 2400000


def test_cells_read_as_written():
    problems = []
    assert batch.cell("2,500,000", "x", problems, True) == 2500000.0
    assert batch.cell(3.0, "x", problems, False) == "3" and batch.cell(" ", "x", problems, True) is None
    assert batch.cell("2.500.000", "here", problems, True) is None and problems == ["here: '2.500.000' is not a number"]
    assert batch.locate(["bids", 1, "lines", 0, "unit_price"], {("bids", 1, "lines", 0): "row 5"}) == "row 5, unit_price"
    assert batch.locate([], {}) == "the package"


def test_analyse_once_per_item_reuses_finished_lines_and_isolates_failures(capsys):
    lines = [{"line": "1", "item": "Toyota Hilux double cabin pickup", "quantity": 1, "prices": {"A": 95e6}},
             {"line": "2", "item": "Toyota Hilux double cabin pickup", "quantity": 1, "prices": {"A": 96e6}},
             {"line": "3", "item": "Laptop computer", "quantity": 1},
             {"line": "4", "item": "Laptop computer", "quantity": 1, "prices": {"A": 2e6}}]
    base = {"currency": "NGN", "year": 2025, "country": "NG"}

    run = db.start_run("package", "analyse", base, lines, ["A"])

    async def go():
        async with fake_session(ScriptedModel(keywords={"Toyota Hilux double cabin pickup": ["hilux"]})) as (graph, _):
            first = await batch.analyse(graph, lines, base, run)
            again = await batch.analyse(graph, lines, base, run, only=["4"])

            class Broken:
                async def aget_state(self, config):
                    return SimpleNamespace(values={})

                async def ainvoke(self, *a, **k):
                    raise RuntimeError("model went away")
            failed = await batch.analyse(Broken(), lines[:1], base, db.start_run("package", "x", base, lines, ["A"]))
            return first, again, failed
    first, again, failed = asyncio.run(go())
    assert first["1"] is first["2"] and "3" not in first and first["4"]["final"]["band"] is None
    assert "reusing the finished result" in capsys.readouterr().out and set(again) == {"4"}
    assert failed["1"]["error"].startswith("RuntimeError")
    replayed = server.review(run)[2]  # the report is rebuilt from the run's recorded steps
    assert replayed["1"]["final"] == first["1"]["final"] and replayed["2"] is replayed["1"] and "3" not in replayed


def test_full_package_run_with_revise_and_approve(session, answers, capsys, monkeypatch):
    args = argparse.Namespace(sheet=str(EXAMPLES / "bid-package.csv"), resume=None, currency="NGN", year=2025, country="NG")
    answers(["revise", "", "check line 1 again", "approve", "Ada Obi"])
    asyncio.run(batch.run(args))
    out = capsys.readouterr().out
    assert "Read bid-package.csv: BoQ with 3 lines, 3 bidders" in out and "Problems in the bid files: 4" in out
    assert '"saved"' in out and "bid-package" in out
    answers(["reject"])
    asyncio.run(batch.run(args))
    out = capsys.readouterr().out
    assert "Not saved: the reviewer rejected the draft." in out
    run = re.search(r"Review (\w+) \|", out).group(1)
    answers(["reject"])
    asyncio.run(batch.run(argparse.Namespace(**{**vars(args), "resume": run})))
    out = capsys.readouterr().out
    assert f"Review {run} |" in out and "reusing the finished result" in out
    with pytest.raises(SystemExit, match="Country ZZ is not loaded"):
        asyncio.run(batch.run(argparse.Namespace(**{**vars(args), "country": "zz", "currency": None})))
    bad = argparse.Namespace(**{**vars(args), "sheet": str(EXAMPLES / "missing.txt")})
    with pytest.raises(SystemExit, match="was not read"):
        asyncio.run(batch.run(bad))


def test_the_documents_checklist_is_read_strictly(tmp_path):
    import datetime as dt
    package, where = batch.load(EXAMPLES / "bid-package.xlsx")
    assert batch.requirements(package)[2] == {"ref": "3", "document": "Tax compliance certificate",
                                              "rule": "Valid on 2025-10-15"}
    found = {(c["check"], c["bidder"], c["data"]["ref"]) for c in batch.document_checks(package, where)}
    assert found == {("document_missing", "Bidder C", "2"), ("document_expired", "Bidder B", "3"),
                     ("document_years", "Bidder C", "4")}
    assert "4 required documents" in batch.describe("x.xlsx", *batch.unpack(package, where), package["documents"])
    wb = load_workbook(EXAMPLES / "bid-package.xlsx")
    wb["Documents"]["E4"] = dt.datetime(2026, 1, 31)  # an Excel date cell reads as a date
    wb.save(tmp_path / "dates.xlsx")
    package, where = batch.load(tmp_path / "dates.xlsx")
    assert package["documents"][2]["bidders"]["Bidder B"] == "2026-01-31"
    assert not [c for c in batch.document_checks(package, where) if c["check"] == "document_expired"]
    for cell, value, expected in (("A1", "Number", "the first columns must be Ref, Document, Rule"),
                                  ("C3", "Valid soon", "rule 'Valid soon' is not Submitted"),
                                  ("D3", "Maybe", "sheet “Documents”, row 3, Bidder A: 'Maybe'; write Yes or No"),
                                  ("D4", None, "row 4, Bidder A: ''; write the date it is valid until"),
                                  ("F1", "Bidder Z", "Bidder Z is not a bidder in this package"),
                                  ("F1", None, "a bidder column with no name")):
        wb = load_workbook(EXAMPLES / "bid-package.xlsx")
        wb["Documents"][cell] = value
        wb.save(tmp_path / "bad.xlsx")
        assert expected in refused(tmp_path / "bad.xlsx")
    wb = load_workbook(EXAMPLES / "bid-package.xlsx")
    for row in list(wb["Documents"].iter_rows()):
        for c in row:
            c.value = None
    wb.save(tmp_path / "empty.xlsx")
    assert "sheet “Documents”: no rows" in refused(tmp_path / "empty.xlsx")
    pkg = json.loads((EXAMPLES / "bid-package.json").read_text())
    pkg["documents"].append(dict(pkg["documents"][0]))
    with pytest.raises(batch.PackageError, match="document ref 1 appears more than once"):
        batch.validate(pkg)
    pkg["documents"] = [{**pkg["documents"][0], "bidders": {"Bidder A": "Yes"}}]
    assert "no column for Bidder B, Bidder C" in "\n".join(pytest.raises(batch.PackageError, batch.validate, pkg).value.problems)
    assert batch.rule_of("Valid on 2025-02-30") is None and batch.rule_of("covers 2023") == ("covers", {2023})
    assert batch.value_of("covers", "Yes") is None and batch.value_of("valid", "31/12/2025") is None


def test_template_in_each_format(tmp_path, monkeypatch):
    for ext in ("xlsx", "csv", "json"):
        said = batch.template(tmp_path / f"t.{ext}")
        assert said.startswith(f"Wrote {tmp_path / f't.{ext}'}") and f"--package {tmp_path / f't.{ext}'}" in said
    assert load_workbook(tmp_path / "t.xlsx").sheetnames == ["BoQ", "Documents", "Bidder 1", "Bidder 2"]
    assert (tmp_path / "t.csv").read_text().startswith("Bidder,Item No,Description,Unit,Quantity,Budget per unit")
    assert set(json.loads((tmp_path / "t.json").read_text())) == {"boq", "documents", "bids"}
    with pytest.raises(SystemExit, match="already exists"):
        batch.template(tmp_path / "t.xlsx")
    with pytest.raises(SystemExit, match="use .xlsx, .csv or .json"):
        batch.template(tmp_path / "t.txt")
    from fairprice import agent
    monkeypatch.setattr("sys.argv", ["fairprice", "--template", str(tmp_path / "u.json")])
    agent.main()
    assert (tmp_path / "u.json").exists()
