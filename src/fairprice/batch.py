"""Review a bid package: every bidder's price schedule for one procurement, line by line.

Usage:
  uv run fairprice --package examples/bid-package.xlsx --year 2025     (or .csv, or .json)
  or, from any MCP client, the start_package_review tool with the package as JSON

A package must match bid_package.schema.json exactly; one that does not is refused with every problem and where it
is, never interpreted. Spreadsheets use the column headings of a PPRA (Kenya) price schedule; HEADINGS lists the
ones accepted for each field. The three forms of the same package:
  xlsx  a sheet named "BoQ" (Item No, Description, Unit, Quantity, and optionally Budget per unit) and one sheet per
        bidder, named after the bidder (Item No, Unit Price, and optionally Description, Unit, Quantity, Total Price,
        Country of origin)
  csv   one file with a Bidder column and the same headings; rows whose bidder is "BoQ" are the BoQ
  json  {"boq": [{line, item, unit, quantity, budget}],
         "bids": [{bidder, lines: [{line, unit_price, item, unit, quantity, total, country_of_origin}]}]}
The agent builds one price band per distinct item, then code compares every bidder's price with it and checks the
bid files. Bidders are never totalled, ranked or recommended.
"""

import csv
import json
import re
import statistics
from itertools import combinations
from pathlib import Path

import jsonschema
from openpyxl import load_workbook

from fairprice import agent, report, server
from fairprice.db import get_run, start_run

MIN_IDENTICAL = 2  # ponytail: fixed threshold for the identical-prices pattern; tune with committee feedback
SCHEMA = json.loads(Path(__file__).with_name("bid_package.schema.json").read_text())
BOQ_COLUMNS = ("line", "item", "unit", "quantity", "budget")
BID_COLUMNS = ("line", "item", "unit", "quantity", "unit_price", "total", "country_of_origin")
CSV_COLUMNS = ("bidder", *dict.fromkeys(BOQ_COLUMNS + BID_COLUMNS))
NUMERIC = {"quantity", "unit_price", "total", "budget"}
HEADINGS = {  # field: the spreadsheet headings for it (PPRA price schedules), compared after plain()
    "bidder": ("Bidder", "Tenderer"),
    "line": ("Item No", "Item No.", "Item Number"),
    "item": ("Description", "Description of Goods"),
    "unit": ("Unit", "Unit of Issue", "Unit of Measure"),
    "quantity": ("Quantity", "Qty"),
    "unit_price": ("Unit Price", "Rate", "Unit Rate"),
    "total": ("Total Price", "Total Amount", "Total"),
    "country_of_origin": ("Country of Origin",),
    "budget": ("Budget per unit", "Budget unit price"),
}


def plain(heading):
    """A heading without case, punctuation at the end, notes in brackets or a currency: "Rate in KES (VAT
    inclusive)" -> "rate"."""
    h = re.sub(r"\(.*?\)", " ", str(heading).lower())
    h = re.sub(r"\bin [a-z]{3}\b", " ", h)
    return " ".join(h.split()).rstrip(".:")


FIELD = {plain(h): field for field, hs in HEADINGS.items() for h in hs}
NUMBER = re.compile(r"^-?(\d+|\d{1,3}(,\d{3})+)(\.\d+)?$")  # 2500000, 2,500,000 or 2500000.50; nothing else


class PackageError(ValueError):
    """A bid package that does not match the schema. Lists every problem and where it is."""

    def __init__(self, path, problems):
        self.problems = problems
        super().__init__(f"{Path(path).name} was not read: it does not match the bid package format "
                         f"(bid_package.schema.json).\n" + "\n".join(f"  - {p}" for p in problems))


def cell(value, where, problems, numeric):
    """A cell as written: text, or a number. Anything that is not plainly a number is a problem, not a guess."""
    if value is None or str(value).strip() == "":
        return None
    if not numeric:
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if NUMBER.match(text := str(value).strip()):
        return float(text.replace(",", ""))
    problems.append(f"{where}: {text!r} is not a number")
    return None


def records(rows, allowed, required, label, problems):
    """[(where, {column: value})] from a table whose first row holds exactly the allowed column names."""
    rows = [list(r) for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        problems.append(f"{label}: no rows")
        return []
    written = [str(c).strip() if c not in (None, "") else "" for c in rows[0]]
    header = [FIELD.get(plain(h), "?") if h else "" for h in written]
    if unknown := [w for w, h in zip(written, header) if h not in allowed and h]:
        problems.append(f"{label}, row 1: unknown column(s) {', '.join(unknown)}; the columns are "
                        f"{', '.join(HEADINGS[c][0] for c in allowed)}")
    if twice := sorted({h for h in header if h and h != "?" and header.count(h) > 1}):
        problems.append(f"{label}, row 1: more than one column for {', '.join(HEADINGS[c][0] for c in twice)}")
    if missing := [c for c in required if c not in header]:
        problems.append(f"{label}, row 1: missing column(s) {', '.join(HEADINGS[c][0] for c in missing)}")
    if unknown or missing or twice:
        return []
    out = []
    for n, r in enumerate(rows[1:], start=2):
        where = f"{label}, row {n}"
        if any(v not in (None, "") for h, v in zip(header, r) if not h) or len(r) > len(header):
            problems.append(f"{where}: a value in a column with no name")
        values = {h: cell(v, f"{where}, {w}", problems, h in NUMERIC) for w, h, v in zip(written, header, r) if h}
        out.append((where, {k: v for k, v in values.items() if v is not None}))
    return out


def read_xlsx(path, problems):
    wb = load_workbook(path, read_only=True, data_only=True)
    package, where = {"boq": [], "bids": []}, {}
    sheets = {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}
    if not any(name.strip().lower() == "boq" for name in sheets):
        problems.append('no sheet named "BoQ" (the bill of quantities)')
    for name, rows in sheets.items():
        label = f"sheet “{name}”"
        if name.strip().lower() == "boq":
            for w, r in records(rows, BOQ_COLUMNS, ("line", "item", "quantity"), label, problems):
                where[("boq", len(package["boq"]))] = w
                package["boq"].append(r)
        else:
            bid = {"bidder": name.strip(), "lines": []}
            for w, r in records(rows, BID_COLUMNS, ("line", "unit_price"), label, problems):
                where[("bids", len(package["bids"]), "lines", len(bid["lines"]))] = w
                bid["lines"].append(r)
            package["bids"].append(bid)
    return package, where


def read_csv(path, problems):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    package, where, bids = {"boq": [], "bids": []}, {}, {}
    for w, r in records(rows, CSV_COLUMNS, ("bidder", "line"), Path(path).name, problems):
        bidder = r.pop("bidder", None)
        if bidder is None:
            problems.append(f"{w}: no bidder (write BoQ for bill-of-quantities rows)")
        elif bidder.lower() == "boq":
            if extra := sorted(set(r) - set(BOQ_COLUMNS)):
                problems.append(f"{w}: a BoQ row has no {', '.join(HEADINGS[c][0] for c in extra)}")
            where[("boq", len(package["boq"]))] = w
            package["boq"].append({k: v for k, v in r.items() if k in BOQ_COLUMNS})
        else:
            if bidder not in bids:
                bids[bidder] = {"bidder": bidder, "lines": []}
                package["bids"].append(bids[bidder])
            j = package["bids"].index(bids[bidder])
            where[("bids", j, "lines", len(bids[bidder]["lines"]))] = w
            bids[bidder]["lines"].append(r)
    if not package["boq"] and not problems:
        problems.append("no BoQ rows (rows whose bidder is BoQ)")
    return package, where


def read_json(path, problems):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8-sig")), {}
    except json.JSONDecodeError as e:
        problems.append(f"not valid JSON: {e}")
        return {}, {}


def locate(path, where):
    """The spreadsheet location of a schema error, or its place in the JSON (bids[1].lines[0].unit_price)."""
    for n in range(len(path), 0, -1):
        if tuple(path[:n]) in where:
            return where[tuple(path[:n])] + "".join(f", {p}" for p in path[n:])
    return "".join(f"[{p}]" if isinstance(p, int) else f".{p}" for p in path).lstrip(".") or "the package"


def validate(package, where=None, label="package"):
    """Check a package (the schema's form) against the schema and the rules it cannot express; raise PackageError
    with every problem and where it is. Returns the package with line numbers as text."""
    where = where or {}
    problems = [f"{locate(list(e.absolute_path), where)}: {e.message}"
                for e in jsonschema.Draft202012Validator(SCHEMA).iter_errors(package)]
    if not problems:  # rules the schema cannot express
        lines = [str(r["line"]) for r in package["boq"]]
        problems += [f"BoQ line {n} appears more than once" for n in sorted({n for n in lines if lines.count(n) > 1})]
        names = [b["bidder"] for b in package["bids"]]
        problems += [f"bidder {n!r} appears more than once" for n in sorted({n for n in names if names.count(n) > 1})]
        for j, bid in enumerate(package["bids"]):
            seen = set()
            for k, r in enumerate(bid["lines"]):
                at, line = locate(["bids", j, "lines", k], where), str(r["line"])
                if line not in lines:
                    problems.append(f"{at}: line {line} is not in the BoQ")
                if line in seen:
                    problems.append(f"{at}: line {line} is priced twice")
                seen.add(line)
    if problems:
        raise PackageError(label, problems)
    for r in package["boq"]:
        r["line"] = str(r["line"])
    for bid in package["bids"]:
        for r in bid["lines"]:
            r["line"] = str(r["line"])
    return package


def load(path):
    """Read a bid package (.xlsx, .csv or .json) into the schema's form, or raise PackageError with every problem.
    Returns (package, where): where maps each row to its location in the file, for findings to cite."""
    readers = {".xlsx": read_xlsx, ".csv": read_csv, ".json": read_json}
    if Path(path).suffix.lower() not in readers:
        raise PackageError(path, [f"use .xlsx, .csv or .json, not {Path(path).suffix or 'no extension'}"])
    if not Path(path).is_file():
        raise PackageError(path, ["file not found"])
    problems = []
    package, where = readers[Path(path).suffix.lower()](path, problems)
    if problems:
        raise PackageError(path, problems)
    return validate(package, where, path), where


def unpack(package, where=None):
    """(boq rows, {bidder: rows}) for the checks and the agent; each row carries where it is in the file."""
    where = where or {}
    boq = [{"where": locate(["boq", i], where), "line": r["line"], "item": r["item"], "quantity": r["quantity"],
            "unit": r.get("unit"), "budget": r.get("budget")} for i, r in enumerate(package["boq"])]
    bidders = {bid["bidder"]: [{"where": locate(["bids", j, "lines", k], where), "line": r["line"],
                                "item": r.get("item", ""), "quantity": r.get("quantity"),
                                "unit_price": r["unit_price"], "total": r.get("total")}
                               for k, r in enumerate(bid["lines"])]
               for j, bid in enumerate(package["bids"])}
    return boq, bidders


def template_bytes(kind):
    """A blank bid package (".xlsx", ".csv" or ".json") to fill in."""
    boq = [HEADINGS[c][0] for c in BOQ_COLUMNS]
    bid = [HEADINGS[c][0] for c in BID_COLUMNS if c != "country_of_origin"]
    if kind == ".json":
        return (json.dumps({"boq": [{"line": "1", "item": "", "unit": "", "quantity": 1, "budget": 0}],
                            "bids": [{"bidder": "Bidder 1", "lines": [{"line": "1", "unit_price": 0}]}]},
                           indent=2) + "\n").encode()
    if kind == ".csv":
        return (",".join([HEADINGS["bidder"][0], *dict.fromkeys(boq + bid)]) + "\n").encode()
    import io
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "BoQ"
    wb["BoQ"].append(boq)
    for name in ("Bidder 1", "Bidder 2"):
        wb.create_sheet(name).append(bid)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template(path):
    """Write a blank bid package to fill in; returns what to do next. Never overwrites a file."""
    path, kind = Path(path), Path(path).suffix.lower()
    if kind not in (".xlsx", ".csv", ".json"):
        raise SystemExit(f"use .xlsx, .csv or .json, not {kind or 'no extension'}")
    if path.exists():
        raise SystemExit(f"{path} already exists; choose another name")
    path.write_bytes(template_bytes(kind))
    return (f"Wrote {path}. Fill in the BoQ and each bidder's price schedule"
            + (" (one sheet per bidder, named after the bidder; rename or add sheets)" if kind == ".xlsx" else "")
            + f". Then: uv run fairprice --package {path}")


def read_package(path):
    return unpack(*load(path))


def describe(path, boq, bidders):
    """What was read, so a misread file is caught before the agent spends minutes on it."""
    out = [f"Read {Path(path).name}: BoQ with {len(boq)} line{'s' if len(boq) != 1 else ''}, {len(bidders)} bidders"]
    for name, rows in bidders.items():
        missing = sorted({r["line"] for r in boq} - {r["line"] for r in rows}, key=lambda x: (len(x), x))
        out.append(f"  {name}: {len(rows)} of {len(boq)} lines priced"
                   + (f" (not priced: {', '.join(missing)})" if missing else ""))
    return "\n".join(out)


def package_lines(boq, bidders):
    """One entry per BoQ line, with every bidder's quoted unit price."""
    lines = {r["line"]: {"line": r["line"], "item": r["item"], "quantity": r["quantity"],
                         **({"budget": r["budget"]} if r.get("budget") is not None else {})} for r in boq}
    for name, rows in bidders.items():
        for r in rows:
            lines[r["line"]].setdefault("prices", {})[name] = r["unit_price"]
    return list(lines.values())


def file_checks(boq, bidders):
    """Problems a committee checks in the bid files themselves. Each one names where in the file it is."""
    found = []

    def flag(check, bidder, r, line=None, **data):
        found.append({"check": check, "bidder": bidder, "line": r["line"] if r else line,
                      "where": r["where"] if r else None, "data": data})

    boq_by_line = {r["line"]: r for r in boq}
    for name, rows in bidders.items():
        for r in rows:
            q, u, t = r["quantity"] or boq_by_line[r["line"]]["quantity"], r["unit_price"], r["total"]
            if q and t is not None and abs(q * u - t) > max(1, 0.001 * abs(t)):
                flag("arithmetic_error", name, r, quantity=q, unit_price=u, total=t)
            expected = boq_by_line[r["line"]]["quantity"]
            if r["quantity"] is not None and r["quantity"] != expected:
                flag("quantity_changed", name, r, quoted=r["quantity"], requested=expected)
        for line in sorted(boq_by_line.keys() - {r["line"] for r in rows}, key=lambda x: (len(x), x)):
            flag("missing_line", name, None, line, item=boq_by_line[line]["item"])
    prices = {name: {r["line"]: r["unit_price"] for r in rows} for name, rows in bidders.items()}
    for a, b in combinations(prices, 2):
        shared = sorted(prices[a].keys() & prices[b].keys(), key=lambda x: (len(x), x))
        same = [line for line in shared if prices[a][line] == prices[b][line]]
        if len(same) >= MIN_IDENTICAL and len(same) * 2 >= len(shared):
            flag("identical_prices", f"{a}, {b}", None, ", ".join(same), same=len(same), shared=len(shared))
    return found


async def analyse(graph, lines, base, run_id, feedback="", attempt=0, only=None, usage=None):
    """Run the agent (up to the draft, no gate) once per distinct item of the run; return {line: state}.
    Every search and check is recorded under the run, line by line."""
    states, by_item = {}, {}
    agent.CURRENT_RUN = run_id
    for n, entry in enumerate(lines):
        if only is not None and entry["line"] not in only:
            continue
        if not entry.get("prices"):
            continue
        key = entry["item"].lower().strip()
        if key in by_item:  # the same item on another line: one band serves both
            states[entry["line"]] = by_item[key]
            continue
        thread = f"{run_id}-{n}" + (f"-r{attempt}" if attempt else "")
        config = {"configurable": {"thread_id": thread}, "callbacks": [usage] if usage else []}
        req = {**base, "run_id": run_id, "line": entry["line"], "item": entry["item"],
               "unit_price": statistics.median(entry["prices"].values())}
        snapshot = await graph.aget_state(config)
        if snapshot.values.get("final") is not None:
            print(f"Line {entry['line']}: reusing the finished result")
        else:
            print(f"\nLine {entry['line']}: {entry['item']}", flush=True)
            try:
                await graph.ainvoke({"request": req, "feedback": feedback}, config, interrupt_after=["recheck"])
            except Exception as e:  # one failed line must not stop the batch
                agent.audit(run_id, event="line_failed", line=entry["line"], error=repr(e))
                print(f"  ! line {entry['line']} failed: {e!r}", flush=True)
                by_item[key] = states[entry["line"]] = {"request": req, "error": repr(e)[:300]}
                continue
        by_item[key] = states[entry["line"]] = (await graph.aget_state(config)).values
    return states


async def run(args):
    try:
        package, where = load(args.sheet)
    except PackageError as e:
        raise SystemExit(str(e))
    boq, bidders = unpack(package, where)
    print(describe(args.sheet, boq, bidders))
    lines = package_lines(boq, bidders)
    checks = file_checks(boq, bidders)  # each with its sheet and row
    names = list(bidders)
    loaded = server.loaded_countries()
    if not args.resume:
        args.country = agent.pick_country(args.country, loaded)
    base = server.review_request(args.currency, args.year, args.country)
    if base is None:
        raise SystemExit(server.not_loaded(args.country))
    run_id = args.resume if args.resume and get_run(args.resume) else start_run(
        "package", Path(args.sheet).name, base, lines, names, checks)
    print(f"Review {run_id} | model {agent.MODEL} | {len(lines)} lines, {len(bidders)} bidders\n"
          f"{agent.country_line(base['country'], base['currency'], loaded)}")
    agent.audit(run_id, event="run_started", sheet=str(args.sheet), bidders=names, lines=len(lines),
                file_checks=checks, model=agent.MODEL)

    def which_lines():
        redo = input("Which line numbers? (comma-separated, blank for all): ").replace(" ", "")
        return redo.split(",") if redo else None
    async with agent.session() as (graph, tools):
        usage = agent.Usage()
        states = await analyse(graph, lines, base, run_id, usage=usage)
        attempt = 0
        while True:
            decision = await agent.gate(run_id, report.summary(lines, states, names, checks, base), usage,
                                        ask_lines=which_lines)
            if decision["action"] != "revise":
                break
            attempt += 1
            states |= await analyse(graph, lines, base, run_id, decision["feedback"], attempt, decision["lines"], usage)
        saved = None
        if decision["action"] == "approve":
            result = await tools["write_price_review"].ainvoke({"run_id": run_id, "approved_by": decision["name"]})
            saved = agent.tool_text(result)
        print(saved or "Not saved: the reviewer rejected it.")
        agent.audit(run_id, event="run_finished", saved=saved, cost=usage.summary())
