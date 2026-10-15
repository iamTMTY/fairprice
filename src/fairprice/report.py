"""The committee-facing report, built from tool outputs only: first as data (stored in the reports table), then
as files in one format: xlsx (default; Prices and Sources sheets), csv (two files) or json (one file).

Written for evaluation committee members, not engineers: plain words and numbered references to the
official published records. Internal record ids appear only as "FairPrice audit ID", to link to the audit log.
"""

import csv
import datetime as dt
import io
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from fairprice.server import assess, rows

DRAFT_STATUS = "Not approved"  # a report built but not written; a written one says "Approved by <name> on <date>"
NOTICE = ("Evidence prepared by FairPrice for the evaluation committee. It does not rank bidders or decide anything; "
          "the committee decides.")

SYMBOL = {"NGN": "₦", "KES": "KSh "}
STATUS = {  # code: (label, icon, fill colour)
    "within_band": ("In line with past prices", "✓", "DFF2DF"),
    "above_band": ("Well above past prices", "▲", "F8D7D7"),
    "below_band": ("Well below past prices", "▼", "FDEFC8"),
    "no_comparison": ("No comparison possible", "–", None),
    "not_priced": ("Not priced", "", None),
    "failed": ("Check did not finish", "!", None),
}
FLAG_TEXT = {
    "multi_item_contract": "the contract covers several different items, so it has no price per item",
    "no_unit_price": "no quantity was published, so there is no price per item",
    "no_quantity_published": "no quantity was published, so there is no price per item",
    "duplicate of an earlier record": "a duplicate of another record of the same purchase",
    "test_record": "a test entry the portal published by mistake (buyer “Test Ministry”)",
    "value_zero_or_negative": "the published value is zero or negative, an entry error",
    "price_below_one": "the price per item is below 1, almost certainly an entry error",
}


def money(x, cur, short=False):
    if x is None:
        return ""
    sym = SYMBOL.get(cur, f"{cur} ")
    if not short:
        return f"{sym}{x:,.0f}"
    for div, unit in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
        if x >= div:
            v = x / div
            digits = 0 if v >= 100 else 1 if v >= 10 else 2
            text = f"{v:.{digits}f}"
            return sym + (text.rstrip("0").rstrip(".") if "." in text else text) + unit
    return f"{sym}{x:,.0f}"


def plain_reason(reason):
    """A rejection reason from the model or a tool flag, in words a committee member can read."""
    if reason.startswith("Challenger: "):
        return "Second check: " + reason.removeprefix("Challenger: ")
    parts = []
    for part in reason.split(", "):
        if part.startswith("extreme_value:"):
            parts.append(f"an extreme value: {part.split(':')[1]}× the median of the other comparable records "
                         "(a typing error, or a total entered as a price per item)")
        elif part.startswith("invalid_date:"):
            year = part.split(":")[1]
            parts.append("the date is missing or a placeholder" + (f" ({year})" if year != "None" else ""))
        else:
            parts.append(FLAG_TEXT.get(part, part))
    return "; ".join(parts)


def ranges(numbers):
    """[1, 2, 3, 5] -> '1–3, 5'."""
    out, nums = [], sorted(numbers)
    for n in nums:
        if out and n == out[-1][1] + 1:
            out[-1][1] = n
        else:
            out.append([n, n])
    return ", ".join(f"{a}–{b}" if b > a else str(a) for a, b in out)


def single(state):
    """A single-item run in the shape of a one-line package, so both use the same report."""
    req = state["request"]
    return ([{"line": "1", "item": req["item"], "quantity": None, "prices": {"Price under review": req["unit_price"]}}],
            {"1": state}, ["Price under review"])


def cells(lines, states, bidders):
    """{(line, bidder): {price, status, ratio}}; the one place a quoted price is judged against its band."""
    out = {}
    for entry in lines:
        st = states.get(entry["line"]) or {}
        band = (st.get("final") or {}).get("band")
        for b in bidders:
            price = (entry.get("prices") or {}).get(b)
            status, ratio = ("not_priced" if price is None else "failed" if not st or st.get("error")
                             else "no_comparison" if not band else None), None
            if status is None:
                ratio, a = assess(price, band["median"])
                status = a.split(":")[0]
            budget = entry.get("budget")
            out[(entry["line"], b)] = {"price": price, "status": status, "ratio": ratio,
                                       "budget": budget, "times_budget": round(price / budget, 2)
                                       if price is not None and budget else None}
    return out


def sources(lines, states):
    """Number every record the report relies on, in order of appearance: ({id: n}, {id: record})."""
    order = []
    for entry in lines:
        final = (states.get(entry["line"]) or {}).get("final") or {}
        for r in final.get("accepted", []) + final.get("excluded", []):
            if r["id"] not in order:
                order.append(r["id"])
    if not order:
        return {}, {}
    info = {r["id"]: r for r in rows(
        "SELECT l.id, l.ocid, l.line_id, l.country, l.year, l.buyer, l.description, l.quantity, l.unit_price, "
        f"l.currency, r.source, s.title AS publisher FROM lines l JOIN releases r USING (country, release_id) "
        f"LEFT JOIN sources s ON s.url = r.source "
        f"WHERE l.id IN ({','.join('?' * len(order))})", order)}
    return {i: n for n, i in enumerate(order, 1)}, info


def check_text(c, cur):
    """A bid file problem as plain sentences, naming where in the file it came from."""
    where = f"{c['bidder']}, line {c['line']}" + (f" ({c['where']})" if c.get("where") else "")
    d = c.get("data") or {}
    if c["check"] == "arithmetic_error":
        return (f"{where}: the line total does not add up. {d['quantity']:g} × {money(d['unit_price'], cur)} = "
                f"{money(d['quantity'] * d['unit_price'], cur)}, but the bid says {money(d['total'], cur)}.")
    if c["check"] == "quantity_changed":
        return f"{where}: priced for {d['quoted']:g}, but the bill of quantities asks for {d['requested']:g}."
    if c["check"] == "missing_line":
        return f"{c['bidder']} did not price line {c['line']} ({d['item']})."
    if c["check"] == "identical_prices":
        return (f"{c['bidder'].replace(', ', ' and ')} quoted exactly the same price per item on {d['same']} of the "
                f"{d['shared']} lines they both priced (lines {c['line']}). This can be coincidence, for example a "
                "published price list; the committee may wish to look into it.")
    return f"{where}: {c['check']}."


def file_name(item, run_id):
    """The report's name in the reports table; its files add an extension (and -prices/-sources for csv)."""
    return f"{re.sub(r'[^a-z0-9]+', '-', item.lower()).strip('-')[:40]}-{run_id}"


def summary(lines, states, bidders, checks, base):
    """A short plain-text summary for the terminal."""
    cur, cs, out = base["currency"], cells(lines, states, bidders), []
    for entry in lines:
        band = ((states.get(entry["line"]) or {}).get("final") or {}).get("band")
        out.append(f"Line {entry['line']} · {entry['item']}: " + (
            f"typical past price {money(band['median'], cur, True)} ({band['n']} purchases)" if band else "no comparison possible"))
        for b in bidders:
            c = cs[(entry["line"], b)]
            label, icon, _ = STATUS[c["status"]]
            ratio = f" ({c['ratio']:g}× typical)" if c["status"] in ("above_band", "below_band") else ""
            over = (f"; {c['times_budget']:g}× the budget" if c["times_budget"] and c["times_budget"] > 1 else "")
            out.append(f"    {b:<20} {money(c['price'], cur, True):>10}  {icon} {label}{ratio}{over}")
    if checks:
        out.append(f"Problems in the bid files: {len(checks)}")
        out += [f"    - {check_text(c, cur)}" for c in checks]
    return "\n".join(out)


FORMATS = ("xlsx", "csv", "json")


def build(title, lines, states, bidders, checks, base, run_id, model=""):
    """The report as data: what the reports table stores, what get_price_review returns, and what every
    file format is written from. Bidders are never totalled or ranked."""
    cur, year = base["currency"], base["year"]
    cs = cells(lines, states, bidders)
    refs, info = sources(lines, states)
    budgets = any(e.get("budget") for e in lines)  # the budget columns only when the BoQ has budgets
    prices = []
    for entry in lines:
        final = (states.get(entry["line"]) or {}).get("final") or {}
        band = final.get("band")
        for b in bidders:
            c = cs[(entry["line"], b)]
            prices.append({
                "line": entry["line"], "item": entry["item"], "quantity_requested": entry["quantity"], "bidder": b,
                "quoted_price": c["price"], "typical_past_price": round(band["median"]) if band else None,
                "times_typical": c["ratio"], "result": STATUS[c["status"]][0],
                **({"budget_per_item": c["budget"], "times_budget": c["times_budget"]} if budgets else {}),
                "past_purchases_compared": band["n"] if band else len(final.get("accepted", [])),
                "sources": sorted(refs[r["id"]] for r in final.get("accepted", [])),
                "problems": [check_text(k, cur) for k in checks or []
                             if b in k["bidder"].split(", ") and entry["line"] in str(k["line"]).split(", ")]})
    adjusted = {r["id"]: r["adjusted_unit_price"] for st in states.values()
                for r in ((st or {}).get("final") or {}).get("accepted", [])}
    srcs = [{"ref": n, "bought_by": (info[i]["buyer"] or "").title() or "(buyer not published)", "year": info[i]["year"],
             "what_was_bought": info[i]["description"], "quantity": info[i]["quantity"],
             "paid_per_item": info[i]["unit_price"], "currency": info[i]["currency"],
             "adjusted_price": round(adjusted[i], 2) if i in adjusted else None,
             "published_by": info[i]["publisher"], "ocid": info[i]["ocid"], "contract_id": info[i]["line_id"],
             "dataset": info[i]["source"], "used": used(i, states), "record_id": i} for i, n in refs.items()]
    how = [{"line": e["line"], "item": e["item"],
            "plan": (states.get(e["line"]) or {}).get("plan"),
            "searches": (states.get(e["line"]) or {}).get("searches", []),
            "assessment": ((states.get(e["line"]) or {}).get("final") or {}).get("assessment"),
            "second_check": {k: v for k, v in ((states.get(e["line"]) or {}).get("challenge") or {}).items()
                             if k in ("verdict", "reason", "applied")} or None,
            "objections": {str(k): v for k, v in (((states.get(e["line"]) or {}).get("challenge") or {}).get("remove") or {}).items()}}
           for e in lines]
    return {"report": {"title": title, "notice": NOTICE, "status": "draft", "approved_by": None, "approved_at": None,
                       "prepared_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
                       "prepared_by": f"FairPrice ({model})" if model else "FairPrice", "run_id": run_id,
                       "currency": cur, "prices_in_year": year, "country": base["country"],
                       "bid_file_checks": [check_text(k, cur) for k in checks] if checks is not None else None},
            "prices": prices, "sources": srcs, "lines": how}


def used(record, states):
    """"Yes", or "No: <why>" for a record a price check left out (a flag, the Analyst's or the second check's reason)."""
    for st in states.values():
        final = (st or {}).get("final") or {}
        if any(r["id"] == record for r in final.get("accepted", [])):
            return "Yes"
    reasons = [e["reason"] for st in states.values() for e in ((st or {}).get("final") or {}).get("excluded", [])
               if e["id"] == record]
    return f"No: {plain_reason(reasons[0])}" if reasons else "No"


def status_line(data):
    r = data["report"]
    if r["status"] != "approved":
        return DRAFT_STATUS
    when = dt.datetime.fromisoformat(r["approved_at"])
    return f"Approved by {r['approved_by']} on {when:%d %B %Y} at {when:%H:%M}"


def headings(data):
    """Plain column names for spreadsheets (xlsx and csv); json keeps the field names."""
    cur, year = data["report"]["currency"], data["report"]["prices_in_year"]
    prices = {"line": "Line", "item": "Item", "quantity_requested": "Quantity requested", "bidder": "Bidder",
              "quoted_price": f"Quoted price per item ({cur})",
              "typical_past_price": f"Typical past price ({cur}, {year} money)",
              "times_typical": "Times the typical price", "result": "Result",
              "budget_per_item": f"Budget per item ({cur})", "times_budget": "Times the budget",
              "past_purchases_compared": "Past purchases compared",
              "sources": "Sources (Ref numbers)", "problems": "Problems in this bid"}
    srcs = {"ref": "Ref", "bought_by": "Bought by", "year": "Year", "what_was_bought": "What was bought",
            "quantity": "Quantity", "paid_per_item": "Paid per item", "currency": "Currency",
            "adjusted_price": f"In {year} money ({cur})", "published_by": "Published by",
            "ocid": "Official record (OCID)", "contract_id": "Contract ID", "dataset": "Dataset",
            "used": "Used in the comparison",
            "record_id": "FairPrice audit ID"}
    flat = lambda k, v: (ranges(v) if k == "sources" else " ".join(v) if k == "problems"
                         else f"L{v}" if k == "record_id"
                         else int(v) if isinstance(v, float) and v.is_integer() else v)  # 2, not 2.0
    return ([{prices[k]: flat(k, v) for k, v in r.items()} for r in data["prices"]],
            [{srcs[k]: flat(k, v) for k, v in r.items()} for r in data["sources"]])


def files(data, name, fmt):
    """{file name: bytes} for the report in one format."""
    if fmt == "json":
        return {f"{name}.json": json.dumps(data, indent=2, ensure_ascii=False).encode()}
    prices, srcs = headings(data)
    if fmt == "csv":
        out = {}
        for part, records in (("prices", prices), ("sources", srcs)):
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(["Status", *(records[0] if records else [])])
            writer.writerows([status_line(data), *r.values()] for r in records)
            out[f"{name}-{part}.csv"] = buf.getvalue().encode("utf-8-sig")  # the BOM lets Excel show ₦ correctly
        return out
    return {f"{name}.xlsx": workbook(data, prices, srcs)}


BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
GREY = PatternFill("solid", fgColor="F0EFEC")


def table(ws, records, widths, status, formats=None):
    """Dict rows as a table with a bold header, then one status line (draft or approved). Returns nothing."""
    for col, key in enumerate(records[0] if records else [], 1):
        cell = ws.cell(1, col, key)
        cell.font, cell.fill, cell.alignment = BOLD, GREY, WRAP
    for row, r in enumerate(records, 2):
        for col, (key, value) in enumerate(r.items(), 1):
            cell = ws.cell(row, col, value)
            cell.alignment = WRAP
            if formats and key in formats:
                cell.number_format = formats[key]
    ws.cell(len(records) + 3, 1, status).font = BOLD  # after a blank row
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.oddHeader.center.text = status
    ws.page_setup.orientation = "landscape"


def workbook(data, prices, srcs):
    """The xlsx file: Prices and Sources sheets."""
    cur, year, status = data["report"]["currency"], data["report"]["prices_in_year"], status_line(data)
    money_fmt = "#,##0"
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    budgets = bool(prices) and f"Budget per item ({cur})" in prices[0]
    table(ws, prices, (6, 30, 11, 18, 16, 18, 11, 24) + ((16, 11) if budgets else ()) + (11, 16, 60), status,
          {f"Quoted price per item ({cur})": money_fmt, f"Typical past price ({cur}, {year} money)": money_fmt,
           "Times the typical price": '0.00"×"', f"Budget per item ({cur})": money_fmt, "Times the budget": '0.00"×"'})
    ws.cell(len(prices) + 4, 1, NOTICE)
    fills = {label: colour for label, _, colour in STATUS.values() if colour}
    for row, r in enumerate(data["prices"], 2):
        if fill := fills.get(r["result"]):
            ws.cell(row, 8).fill = PatternFill("solid", fgColor=fill)
    table(wb.create_sheet("Sources"), srcs, (6, 28, 7, 60, 9, 15, 9, 16, 30, 30, 12, 40, 50, 14), status,
          {"Paid per item": money_fmt, f"In {year} money ({cur})": money_fmt})
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
