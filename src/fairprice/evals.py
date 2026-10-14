"""Run evals/cases.json against the real agent, several times each, and record every result.

Usage: uv run fairprice-evals --repeats 3 [--model ollama:gemma4] [--cases hilux-above,no-benchmark]
Every run is stopped at the approval gate with "reject", which also tests that nothing is saved.
"""

import argparse
import asyncio
import contextlib
import io
import os
import tempfile
import datetime as dt
import json
import re
import statistics
import time
from pathlib import Path
from unittest import mock

from langgraph.types import Command

from fairprice import agent, batch, db, report, server

ROOT = Path(__file__).resolve().parents[2]
CASES = ROOT / "evals" / "cases.json"
RESULTS = ROOT / "evals" / "results"
KNOWN_RECORD = 111893  # a real record id in the default data, for the checks that need one


def judge(case, state):
    final = state.get("final") or {}
    assessment = (final.get("assessment") or "unresolved").split(":")[0]
    accepted = [r["id"] for r in final.get("accepted", [])]
    failures = []
    if assessment not in case["expect"]:
        failures.append(f"assessment {assessment}, expected {'/'.join(case['expect'])}")
    if leaked := sorted(set(accepted) & set(case["traps"])):
        failures.append(f"trap records in band: {leaked}")
    if "max_accepted" in case and len(accepted) > case["max_accepted"]:
        failures.append(f"{len(accepted)} records accepted, expected at most {case['max_accepted']}")
    if state.get("saved"):
        failures.append("report saved although the reviewer rejected it")
    return {"assessment": assessment, "n": len(accepted), "median": (final.get("band") or {}).get("median"),
            "challenge": (state.get("challenge") or {}).get("verdict"), "failures": failures,
            "passed": not failures}


def cost(usage):
    c = usage.summary()
    return {k: c[k] for k in ("model_calls", "input_tokens", "output_tokens", "model_seconds")}


async def run_package(graph, case, model):
    """A whole bid package: per-bidder assessments from the shared bands, judged from the run as recorded
    (a replay of its steps), and nothing saved or ranked."""
    usage = agent.Usage()
    base = server.review_request(case.get("currency"), case["year"], case.get("country", "NG"))
    run_id = None
    start = time.perf_counter()
    try:  # everything, including reading the package: one bad case must not stop the suite
        package, where = batch.load(ROOT / case["package"])
        boq, bidders = batch.unpack(package, where)
        lines = batch.package_lines(boq, bidders)
        checks = batch.file_checks(boq, bidders)
        run_id = db.start_run("package", case["package"], base, lines, list(bidders), checks)
        agent.audit(run_id, event="run_started", sheet=case["package"], model=model)
        await batch.analyse(graph, lines, base, run_id, usage=usage)
        run, lines, states = server.review(run_id)
        got = {f"{line}/{b}": c["status"] for (line, b), c in report.cells(lines, states, run.bidders).items()}
        data = report.build(run.title, lines, states, run.bidders, run.checks, run.request, run_id)
        table = [{**data["prices"][0], **report.headings(data)[0][0]}]  # field names and spreadsheet headings
        failures = [f"line {k}: {got.get(k)}, expected {'/'.join(v)}" for k, v in case["expect_cells"].items()
                    if got.get(k) not in v]
        failures += [f"ranking column {c!r}" for c in table[0] if re.search(r"rank|total|recommend|lowest", c, re.I)]
        result = {"assessment": " ".join(f"{k}={v}" for k, v in got.items()), "n": len(got), "median": None,
                  "challenge": None, "failures": failures, "passed": not failures}
    except Exception as e:
        result = {"assessment": "error", "n": 0, "median": None, "challenge": None,
                  "failures": [repr(e)[:300]], "passed": False}
    return {"case": case["id"], "run_id": run_id, "seconds": round(time.perf_counter() - start, 1), **result,
            **cost(usage)}


async def run_case(graph, tools, case, model):
    if "package" in case:
        return await run_package(graph, case, model)
    usage, run_id = agent.Usage(), None
    start = time.perf_counter()
    try:  # a crash is a failed run, not a reason to stop the suite
        started = agent.as_json(agent.tool_text(await tools["start_review"].ainvoke(
            {"item": case["item"], "unit_price": case["unit_price"], "year": case["year"],
             "currency": case.get("currency", ""), "country": case.get("country", "NG")})))
        run_id = agent.CURRENT_RUN = started["run_id"]
        req = {"run_id": run_id, "line": "1", "item": case["item"], "unit_price": case["unit_price"],
               "currency": started["currency"], "year": started["year"], "country": started["country"]}
        agent.audit(run_id, event="run_started", request=req, model=model)
        config = {"configurable": {"thread_id": run_id}, "callbacks": [usage]}
        await graph.ainvoke({"request": req}, config)
        await graph.ainvoke(Command(resume={"action": "reject"}), config)
        state = (await graph.aget_state(config)).values
        result = judge(case, state)
        final = state.get("final") or {}
        replayed = (server.review(run_id)[2].get("1") or {}).get("final")
        if "accepted" in final and replayed != final:  # a band the tools built (not an unresolved run)
            result["failures"].append("the run's recorded steps do not replay to the same result")
            result["passed"] = False
    except Exception as e:
        result = {"assessment": "error", "n": 0, "median": None, "challenge": None,
                  "failures": [repr(e)[:300]], "passed": False}
    return {"case": case["id"], "run_id": run_id, "seconds": round(time.perf_counter() - start, 1), **result,
            **cost(usage)}


def checked_run(record_ids=None, result=None):
    """A review of one item whose line was searched and checked, without a model. With `result`, the check step
    is written straight into the run (as if the database had been tampered with), skipping the tool's guards."""
    record_ids = [KNOWN_RECORD] if record_ids is None else record_ids
    run_id = server.start_review("test item", 1.0, "NGN", 2025, "NG")["run_id"]
    db.add_step(run_id, "1", "search", {"keywords": ["test"], "country": "NG", "matches": len(record_ids),
                                        "ids": record_ids})
    if result is None:
        server.check_price_comparability(run_id, record_ids)
    else:
        db.add_step(run_id, "1", "check", {"candidate_ids": record_ids, "rejected": {}, "result": result})
    return run_id


@contextlib.contextmanager
def sandbox(provider="file", fmt="xlsx"):
    """A throwaway database and report store (local folder, or a mocked S3 bucket), so checks never touch real reports."""
    from moto import mock_aws

    tmp = tempfile.mkdtemp()
    env = {"FAIRPRICE_SQL_URL": f"sqlite:///{tmp}/check.sqlite", "FAIRPRICE_STORAGE": provider,
           "FAIRPRICE_REPORTS": tmp, "FAIRPRICE_REPORT_FORMAT": fmt, "FAIRPRICE_S3_BUCKET": "fairprice-test",
           "AWS_ACCESS_KEY_ID": "test", "AWS_SECRET_ACCESS_KEY": "test", "AWS_DEFAULT_REGION": "us-east-1"}
    with mock.patch.dict(os.environ, env), mock_aws(), contextlib.redirect_stdout(io.StringIO()):
        if provider == "s3":
            import boto3
            boto3.client("s3").create_bucket(Bucket="fairprice-test")
        yield


def approve(run_id):
    """A person's approval of the review as it is now."""
    return db.grant("write_price_review", server.version(run_id), "A. Officer")


def round_trip(provider, fmt):
    """Start -> search -> check -> draft -> approve -> list -> retrieve -> the files say who approved -> the draft
    files are gone."""
    from fairprice.storage import storage

    with sandbox(provider, fmt):
        run_id = checked_run()
        drafted = server.write_price_review(run_id)
        draft_files = [f"{drafted.get('draft')}.{ext}" for ext in ("xlsx", "csv", "json")]
        saved = server.write_price_review(run_id, approve(run_id))
        listed = server.get_price_review()
        got = server.get_price_review(run_id)
        texts = []
        for name in report.files(got, saved.get("saved", ""), fmt) if "saved" in saved else []:
            content = storage().load(name)
            if fmt == "xlsx":
                from openpyxl import load_workbook
                ws = load_workbook(io.BytesIO(content))["Prices"]
                content = str(ws.cell(ws.max_row, 1).value).encode()
            texts.append(content.decode("utf-8-sig"))
        drafts_left = []
        for name in draft_files:
            try:
                storage().load(name)
                drafts_left.append(name)
            except Exception:
                pass
    passed = (drafted.get("status") == "needs_approval"
              and [r["run_id"] for r in listed.get("reports", [])] == [run_id]
              and got.get("report", {}).get("approved_by") == "A. Officer" and texts and not drafts_left
              and all("A. Officer" in t for t in texts))
    return {"case": f"report ({fmt}, {provider} storage): drafted, approved, listed, retrieved as json, draft removed",
            "passed": bool(passed), "detail": {"saved": saved, "listed": listed, "drafts_left": drafts_left}}


def deterministic_checks():
    """Run and gate checks that need no model: a report is built only from what its run recorded, needs a
    person's approval of that exact version, and is refused when unsourced."""
    empty = {"assessment": "insufficient_comparables", "accepted": [], "excluded": []}
    forged = {"assessment": "within_band", "accepted": [{"id": 999999999}], "excluded": []}
    with sandbox():
        a, b, c, d, e = checked_run(), checked_run(result=empty), checked_run([999999999], forged), checked_run(), checked_run()
        asked = server.write_price_review(a)
        stale = approve(e)
        server.check_price_comparability(e, [KNOWN_RECORD], {str(KNOWN_RECORD): "not comparable"})
        checks = {
            "write refused with an approval for a different review": server.write_price_review(a, approve(d)),
            "write refused with an approval of an older version of the review": server.write_price_review(e, stale),
            "write refused for a review that cites no records": server.write_price_review(b, approve(b)),
            "write refused for a review citing a record that does not exist": server.write_price_review(c, approve(c)),
            "write refused for a review that does not exist": server.write_price_review("../pyproject.toml", "x"),
            "check refused when none of its records came from the review's searches":
                server.check_price_comparability(a, [KNOWN_RECORD + 1]),
            "search refused for a review that does not exist": server.find_comparables("nope", ["laptop"]),
        }
        once = approve(d)
        first, second = (db.use_approval(once, "write_price_review", server.version(d)),
                         db.use_approval(once, "write_price_review", server.version(d)))
        server.discard_review(a)  # the person rejected the draft
        db.decide(asked["request"], True, "A. Officer")
        discarded = server.write_price_review(a, asked["request"])
        replay = server.review(d)[2]["1"]["final"] == server.check_price_comparability(d, [KNOWN_RECORD])
        mixed = server.check_price_comparability(d, [KNOWN_RECORD, KNOWN_RECORD + 1])
    results = [{"case": "write without a person's approval only drafts and records a request", "detail": asked,
                "passed": asked.get("status") == "needs_approval" and "fairprice --approve" in asked.get("message", "")}]
    results += [{"case": name, "passed": "error" in out, "detail": out} for name, out in checks.items()]
    results.append({"case": "a person's approval can be used only once", "passed": bool(first) and second is None,
                    "detail": None})
    results.append({"case": "a rejected draft can no longer be approved", "passed": "error" in discarded,
                    "detail": discarded})
    results.append({"case": "a review rebuilt from its recorded steps matches the tool's answer", "passed": replay,
                    "detail": None})
    results.append({"case": "a record the review's searches did not find is left out of the band, and named",
                    "passed": [r["id"] for r in mixed.get("accepted", [])] == [KNOWN_RECORD]
                    and mixed.get("left_out", {}).get("ids") == [KNOWN_RECORD + 1], "detail": mixed})
    results += [round_trip("file", fmt) for fmt in report.FORMATS] + [round_trip("s3", "xlsx")]
    packages = [batch.load(ROOT / "examples" / f"bid-package.{ext}")[0] for ext in ("xlsx", "csv", "json")]
    results.append({"case": "the xlsx, csv and json templates read to the same bid package",
                    "passed": packages[0] == packages[1] == packages[2], "detail": None})
    head = "Bidder,Item No,Description,Quantity,Unit Price\n"
    bad = {"unknown column": "Bidder,Item No,Description,Quantity,Price per piece\nBoQ,1,Laptop,2,\nBidder A,1,,2,500\n",
           "price in words": head + "BoQ,1,Laptop,2,\nBidder A,1,,2,two million\n",
           "line not in the BoQ": head + "BoQ,1,Laptop,2,\nBidder A,9,,2,500\n",
           "no BoQ": head + "Bidder A,1,,2,500\n"}
    for name, text in bad.items():
        path = Path(tempfile.mkdtemp()) / "package.csv"
        path.write_text(text)
        try:
            batch.load(path)
            refused = None
        except batch.PackageError as e:
            refused = e.problems
        results.append({"case": f"a bid package with {name} is refused, with its location", "passed": bool(refused),
                        "detail": refused})
    package, where = batch.load(ROOT / "examples" / "bid-package.xlsx")
    found = {(c["check"], c["bidder"], c["line"]) for c in batch.file_checks(*batch.unpack(package, where))}
    planted = {("arithmetic_error", "Bidder B", "2"), ("quantity_changed", "Bidder C", "1"),
               ("missing_line", "Bidder C", "3"), ("identical_prices", "Bidder A, Bidder C", "1, 2")}
    results.append({"case": "bid file checks find the 4 planted problems, and only those",
                    "passed": found == planted, "detail": sorted(found ^ planted)})
    return results


async def main_async(args):
    cases = json.loads(CASES.read_text())
    if args.cases:
        cases = [c for c in cases if c["id"] in args.cases.split(",")]
    RESULTS.mkdir(parents=True, exist_ok=True)
    out = RESULTS / f"{dt.datetime.now():%Y%m%d-%H%M}-{args.model.replace(':', '_').replace('/', '_')}.jsonl"
    results = []
    async with agent.session(args.model) as (graph, tools):
        for case in cases:
            for i in range(args.repeats):
                r = await run_case(graph, tools, case, args.model)
                results.append(r)
                with out.open("a") as f:
                    f.write(json.dumps(r) + "\n")
                print(f"{case['id']:<22} run {i + 1}: {'PASS' if r['passed'] else 'FAIL'} "
                      f"{r['assessment']:<26} n={r['n']:<3} {r['seconds']:>6}s  {'; '.join(r['failures'])}", flush=True)
    det = deterministic_checks()
    print("\nSummary", args.model)
    for case in cases:
        rs = [r for r in results if r["case"] == case["id"]]
        medians = [r["median"] for r in rs if r["median"]]
        spread = f"median range {min(medians):,.0f}-{max(medians):,.0f}" if len(medians) > 1 else ""
        print(f"  {case['id']:<22} {sum(r['passed'] for r in rs)}/{len(rs)} passed  "
              f"avg {statistics.mean(r['seconds'] for r in rs):.0f}s  "
              f"avg {statistics.mean(r.get('input_tokens', 0) for r in rs):,.0f} in / "
              f"{statistics.mean(r.get('output_tokens', 0) for r in rs):,.0f} out tokens  {spread}")
    for d in det:
        print(f"  {d['case']:<60} {'PASS' if d['passed'] else 'FAIL'}")
    print(f"\nRaw results: {out}")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--repeats", type=int, default=3)
    p.add_argument("--model", default=agent.MODEL)
    p.add_argument("--cases", help="comma-separated case ids")
    asyncio.run(main_async(p.parse_args()))


if __name__ == "__main__":
    main()
